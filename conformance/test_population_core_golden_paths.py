"""End-to-end golden proofs for the Release B1 population core."""

from __future__ import annotations

from copy import deepcopy
import json
from pathlib import Path
import sys
import tempfile
import unittest


ROOT = Path(__file__).resolve().parents[1]
PANEL_SCRIPTS = ROOT / "skills" / "audience-panel-builder" / "scripts"
AD_TESTING_SCRIPTS = ROOT / "skills" / "audience-ad-testing-lab" / "scripts"
DATA_SCRIPTS = ROOT / "skills" / "audience-data-lab" / "scripts"
PUBLIC_FIXTURES = (
    ROOT / "conformance" / "fixtures" / "population" / "public-proxy"
)
AUTHORIZED_INPUTS = ROOT / "conformance" / "fixtures" / "authorized-audience"
AUTHORIZED_FIXTURES = (
    ROOT / "conformance" / "fixtures" / "population" / "authorized-marketplace"
)
sys.path.insert(0, str(PANEL_SCRIPTS))
sys.path.insert(0, str(AD_TESTING_SCRIPTS))
sys.path.insert(0, str(DATA_SCRIPTS))

from audience_data_lab.authorized_mapping import mapping_sha256  # noqa: E402
from audience_data_lab.authorized_source import (  # noqa: E402
    profile_authorized_bundle,
)
from audience_data_lab.authorized_transform import (  # noqa: E402
    transform_authorized_bundle,
    validate_authorized_handoff,
)
from audience_lab.audience_research_v3 import (  # noqa: E402
    validate_audience_research_v3,
    validate_research_brief_v3,
    validate_saved_panel_v3,
)
from audience_panel_builder.common import (  # noqa: E402
    canonical_json_bytes,
    sha256_json,
    sha256_text,
)
from audience_panel_builder.population.adapters.bls_oews import (  # noqa: E402
    BlsOewsAdapter,
)
from audience_panel_builder.population.adapters.census_cbp import (  # noqa: E402
    CensusCbpAdapter,
)
from audience_panel_builder.population.adapters.census_susb import (  # noqa: E402
    CensusSusbAdapter,
)
from audience_panel_builder.population.adapters.authorized_handoff import (  # noqa: E402
    AuthorizedHandoffAdapter,
)
from audience_panel_builder.population.composition import (  # noqa: E402
    build_composition_plan,
)
from audience_panel_builder.population.feedback import (  # noqa: E402
    bind_outcome_feedback,
    propose_calibration_refresh,
)
from audience_panel_builder.population.frame import (  # noqa: E402
    build_population_frame,
)
from audience_panel_builder.population.validity import (  # noqa: E402
    assess_population_validity,
    finalize_validity_profile,
)
from conformance import (  # noqa: E402
    test_authorized_audience_transform as authorized_transform_test_support,
)


BUILT_AT = "2026-07-24T13:00:00Z"
COMPOSED_AT = "2026-07-23T12:00:00Z"
AUTHORIZED_PROFILED_AT = "2026-07-24T12:00:00Z"
AUTHORIZED_TRANSFORMER_VERSION = "1.0.0"
_BRIEF_V3_KEYS = {
    "panel_tier",
    "evidence_basis",
    "workflow_state_binding",
    "population_frame_result_sha256",
    "population_frame_sha256",
    "authorized_audience_import",
    "structural_findings",
    "overlay_findings",
    "claim_boundary",
    "dimensional_validity",
    "scoped_approvals",
}
_PANEL_V3_KEYS = {
    "panel_tier",
    "evidence_basis",
    "brief_id",
    "population_frame_result_sha256",
    "population_frame_sha256",
    "composition_plan_sha256",
    "validity_profile_sha256",
    "authorized_handoff_sha256",
    "audit_binding",
    "claim_boundary",
    "package_status",
}


def _load(path: Path) -> dict[str, object]:
    return json.loads(path.read_text(encoding="utf-8"))


def _v2_projection(
    document: dict[str, object],
    *,
    brief: bool,
) -> dict[str, object]:
    extension_keys = _BRIEF_V3_KEYS if brief else _PANEL_V3_KEYS
    projection = {
        key: deepcopy(value)
        for key, value in document.items()
        if key not in extension_keys
    }
    projection["schema_version"] = (
        "audience-research-brief-v2"
        if brief
        else "saved-audience-panel-v2"
    )
    return projection


def _authorized_request(
    *,
    calibration_factor: float = 3.0,
    modeled_share: float | None = None,
) -> dict[str, object]:
    modeled_rules = []
    if modeled_share is not None:
        modeled_rules.append({
            "rule_id": "declared-returns-specialists",
            "unit": "eligible-cohort-member",
            "denominator": "all-respondents",
            "dimension_values": {"cohort": "returns_specialists"},
            "method": "declared_weight",
            "structural_weight": modeled_share,
            "uncertainty": {
                "lower": modeled_share,
                "upper": modeled_share,
            },
            "rationale": "Predeclared residual share for downgrade testing.",
        })
    return {
        "schema_version": "audience-frame-request-v1",
        "request_id": "fictional-marketplace-frame",
        "target_audience": "Authorized aggregate marketplace-seller cohort",
        "decision": "Construct an authorized audience-calibrated panel.",
        "desired_claim": "Represent only the exact authorized aggregate cohort.",
        "geography": ["US"],
        "time_basis": {"as_of": "2026-07-24", "lookback_days": 365},
        "target_unit": "eligible-cohort-member",
        "proxy_universes": [{
            "universe_id": "authorized-marketplace-seller-cohort",
            "description": "All eligible members in the authorized aggregate export.",
            "unit": "eligible-cohort-member",
            "denominator": "all-respondents",
            "exact": True,
        }],
        "required_dimensions": ["cohort"],
        "required_joints": [],
        "modeled_cell_rules": modeled_rules,
        "calibration_rules": [{
            "rule_id": "operations-calibration",
            "unit": "eligible-cohort-member",
            "denominator": "all-respondents",
            "dimension_values": {"cohort": "operations_leaders"},
            "calibration_factor": calibration_factor,
            "rationale": "Approved aggregate cohort calibration diagnostic.",
        }],
        "exclusions": [],
        "authorized_evidence_bases": ["first_party_aggregate"],
        "available_capabilities": ["authorized-handoff"],
        "downgrade_policy": {
            "allow_tier_1": True,
            "allow_experimental": True,
            "reason": "Downgrade rather than overclaim modeled structure.",
        },
    }


def _authorized_composition_inputs() -> tuple[
    list[dict[str, object]],
    list[dict[str, object]],
    list[dict[str, object]],
]:
    structural = [{
        "structural_group_id": "operations-sellers",
        "cell_ids": ["operations-leaders"],
        "structural_finding_ids": ["finding-cohort-structure"],
        "evidence_ids": ["evidence-cohort-structure"],
        "must_cover": True,
        "planning_allocation": None,
    }, {
        "structural_group_id": "finance-sellers",
        "cell_ids": ["finance-leaders"],
        "structural_finding_ids": ["finding-cohort-structure"],
        "evidence_ids": ["evidence-cohort-structure"],
        "must_cover": True,
        "planning_allocation": None,
    }]
    overlay_specs = (
        (
            "operations-reliability",
            "Operational reliability and handoff confidence.",
            "finding-operations",
            "evidence-operations",
            "operations",
        ),
        (
            "pricing-confidence",
            "Pricing clarity and cost predictability.",
            "finding-pricing",
            "evidence-pricing",
            "pricing",
        ),
        (
            "returns-friction",
            "Returns handling and recovery friction.",
            "finding-returns",
            "evidence-returns",
            "returns",
        ),
        (
            "customer-service-confidence",
            "Customer-service responsiveness and issue resolution.",
            "finding-customer-service",
            "evidence-customer-service",
            "customer-service",
        ),
    )
    overlays = []
    for (
        overlay_id,
        description,
        finding_id,
        evidence_id,
        topic,
    ) in overlay_specs:
        social_evidence_id = f"social-evidence-{topic}"
        overlays.append({
            "overlay_id": overlay_id,
            "description": description,
            "allocation_basis": "observed",
            "finding_ids": [finding_id],
            "evidence_ids": [evidence_id, social_evidence_id],
            "topic_bindings": [{
                "topic_id": topic,
                "evidence_ids": [social_evidence_id],
            }],
            "decision_relevance": "topic_bound",
        })
    profiles = [{
        "status": "supported",
        "profile_id": "operations-reliability-service",
        "structural_group_id": "operations-sellers",
        "overlay_ids": [
            "operations-reliability",
            "customer-service-confidence",
        ],
        "support_finding_ids": [
            "finding-cohort-structure",
            "finding-operations",
            "finding-customer-service",
        ],
        "support_evidence_ids": [
            "evidence-cohort-structure",
            "evidence-operations",
            "evidence-customer-service",
            "social-evidence-operations",
            "social-evidence-customer-service",
        ],
        "conditional_overlay_allocation": 1.0,
    }, {
        "status": "supported",
        "profile_id": "finance-pricing-returns",
        "structural_group_id": "finance-sellers",
        "overlay_ids": ["pricing-confidence", "returns-friction"],
        "support_finding_ids": [
            "finding-cohort-structure",
            "finding-pricing",
            "finding-returns",
        ],
        "support_evidence_ids": [
            "evidence-cohort-structure",
            "evidence-pricing",
            "evidence-returns",
            "social-evidence-pricing",
            "social-evidence-returns",
        ],
        "conditional_overlay_allocation": 1.0,
    }]
    return structural, overlays, profiles


def _topic_social_evidence() -> dict[str, object]:
    return {
        "records": [{
            "cohort_id": "operations_leaders",
            "evidence_id": "evidence-operations",
            "topic": "operations",
            "summary": "The cohort emphasizes reliable operational handoffs.",
        }, {
            "cohort_id": "finance_leaders",
            "evidence_id": "evidence-pricing",
            "topic": "pricing",
            "summary": "The cohort emphasizes clear pricing and predictable costs.",
        }, {
            "cohort_id": "finance_leaders",
            "evidence_id": "evidence-returns",
            "topic": "returns",
            "summary": "The cohort emphasizes transparent returns handling.",
        }, {
            "cohort_id": "operations_leaders",
            "evidence_id": "evidence-customer-service",
            "topic": "customer_service",
            "summary": "The cohort emphasizes responsive customer service.",
        }],
    }


def _authorized_semantic_inputs() -> dict[str, dict[str, object]]:
    """Return the shape-neutral semantic bytes named by provenance hashes."""

    frame_records = [{
        "cohort_id": "finance-leaders",
        "respondent_count": 680,
        "share": 0.62,
    }, {
        "cohort_id": "operations-leaders",
        "respondent_count": 420,
        "share": 0.38,
    }]
    structured_records = []
    social_records = []
    for record in _topic_social_evidence()["records"]:
        cohort_id = str(record["cohort_id"]).replace("_", "-")
        topic = str(record["topic"]).replace("_", "-")
        structured_records.append({
            "cohort_id": cohort_id,
            "evidence_id": record["evidence_id"],
            "summary": record["summary"],
            "topic": topic,
        })
        social_records.append({
            "cohort_id": cohort_id,
            "observation_id": f"social-evidence-{topic}",
            "summary": record["summary"],
            "topic": topic,
        })
    structured_records.sort(key=lambda item: item["evidence_id"])
    social_records.sort(key=lambda item: item["observation_id"])
    return {
        "frame": {
            "schema_version": "authorized-frame-semantic-input-v1",
            "denominator": "all-respondents",
            "records": frame_records,
            "unit": "eligible-cohort-member",
        },
        "structured": {
            "schema_version": "authorized-structured-semantic-input-v1",
            "records": structured_records,
        },
        "social": {
            "schema_version": "authorized-social-semantic-input-v1",
            "records": social_records,
        },
    }


def _assert_no_repeated_digest_placeholders(documents: object) -> None:
    """Reject all-zero and repeated-character SHA-256 placeholders."""

    def walk(value: object, path: tuple[str, ...]) -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                walk(child, (*path, str(key)))
        elif isinstance(value, list) or isinstance(value, tuple):
            for index, child in enumerate(value):
                walk(child, (*path, f"[{index}]"))
        elif isinstance(value, str):
            digest = value.removeprefix("sha256:")
            if (
                len(digest) == 64
                and len(set(digest)) == 1
                and digest[0] in "0123456789abcdef"
            ):
                raise AssertionError(
                    "repeated-character digest placeholder at "
                    + ".".join(path)
                )

    walk(documents, ())


def _apply_authorized_golden_metadata(
    mapping: dict[str, object],
) -> dict[str, object]:
    semantic_inputs = _authorized_semantic_inputs()
    frame_metadata = mapping["expected_outputs"][0]["metadata"]
    frame_metadata["source"].update({
        "publisher": "Authorized aggregate source",
        "program": "Approved marketplace-seller cohort export",
        "edition": "2026 approved aggregate export",
    })
    frame_metadata["selection_notes"] = (
        "The eligible aggregate cohort was fixed before creative review."
    )
    frame_metadata["coverage_notes"] = (
        "The export covers every eligible member in the authorized cohort."
    )
    frame_metadata["citations"] = [
        "Approved authorized aggregate cohort export, 2026"
    ]
    frame_metadata["raw_snapshot_sha256"] = sha256_json(
        semantic_inputs["frame"]
    )

    topics = (
        "operations",
        "pricing",
        "returns",
        "customer-service",
    )
    for index, operation in enumerate(mapping["operations"]):
        if operation["operation_id"] == "sort-social":
            mapping["operations"][index] = {
                "operation_id": "map-social-evidence-ids",
                "op": "category_map",
                "input": "evidence-final",
                "output": "social-final",
                "field": "evidence_id",
                "mapping": {
                    f"evidence-{topic}": f"social-evidence-{topic}"
                    for topic in topics
                },
                "unmapped": "error",
            }
            break
    else:  # pragma: no cover - protects the imported fixture helper contract.
        raise AssertionError("authorized mapping omitted sort-social")
    structured_metadata = mapping["expected_outputs"][1]["metadata"]
    structured_metadata["input_sha256"] = sha256_json(
        semantic_inputs["structured"]
    )
    structured_summaries = {
        record["evidence_id"]: record["summary"]
        for record in semantic_inputs["structured"]["records"]
    }
    structured_metadata["item_metadata"] = {
        f"evidence-{topic}": {
            "source_url": f"https://example.test/research/{topic}",
            "item_type": "approved-aggregate-affinity-summary",
            "text_fidelity": "faithful-summary",
            "content_sha256": sha256_text(
                structured_summaries[f"evidence-{topic}"]
            ),
            "source_pointer": f"approved-evidence#{topic}",
            "upstream_source_ids": ["authorized-aggregate-cohort-export"],
            "use_constraints": ["qualitative-overlay-only"],
            "quality_flags": [],
        }
        for topic in topics
    }

    social_metadata = mapping["expected_outputs"][2]["metadata"]
    social_metadata["input_sha256"] = sha256_json(
        semantic_inputs["social"]
    )
    social_metadata.update({
        "query": (
            "operations, pricing, returns, and customer-service decision evidence"
        ),
        "source_status": {"authorized_aggregate_source": "ok"},
        "coverage_warnings": [
            "Topic-bound aggregate summaries do not estimate prevalence."
        ],
    })
    social_metadata["collection"].update({
        "provider": "Authorized aggregate source",
        "run_or_dataset_id": "approved-topic-evidence-2026-07-24",
        "collection_method": "permissioned aggregate topic-summary export",
        "access_route": "authorized local normalized export",
        "item_limit": 4,
    })
    social_summaries = {
        record["observation_id"]: record["summary"]
        for record in semantic_inputs["social"]["records"]
    }
    social_metadata["observation_metadata"] = {
        f"social-evidence-{topic}": {
            "platform": "authorized_aggregate_source",
            "source_item_id": f"approved-topic-{topic}",
            "source_url": f"https://example.test/research/{topic}",
            "published_at": None,
            "collected_at": "2026-07-24T12:00:00Z",
            "unit_of_analysis": "approved_aggregate_theme",
            "title": f"{topic.replace('-', ' ').title()} theme",
            "text_fidelity": "faithful_summary",
            "content_sha256": sha256_text(
                social_summaries[f"social-evidence-{topic}"]
            ),
            "engagement": {"prevalence_weight": 0},
            "relevance_score": None,
            "cluster_id": topic,
            "role_status": "unknown",
            "author_group_token": None,
            "freshness_verdict": "approved_export",
            "json_pointer": f"/records/{index}",
            "use_constraints": [
                "qualitative_context_only",
                "no_prevalence_or_weighting",
            ],
            "quality_flags": ["summary_not_verbatim"],
        }
        for index, topic in enumerate(topics)
    }
    outcome_source_sha256 = mapping["input_hashes"][
        "aggregate-outcomes.json"
    ]
    for expected_output in mapping["expected_outputs"]:
        if expected_output["filename"].startswith("outcome-feedback-"):
            expected_output["metadata"]["source_sha256"] = (
                outcome_source_sha256
            )
    mapping["approval"]["mapping_sha256"] = mapping_sha256(mapping)
    return mapping


def _public_batches(
    request: dict[str, object],
) -> list[dict[str, object]]:
    adapters = (
        (
            BlsOewsAdapter(),
            "bls-oews-may-2025.json",
            "bls-public-proxy-frame-batch",
        ),
        (
            CensusSusbAdapter(),
            "census-susb-2022.json",
            "susb-public-proxy-frame-batch",
        ),
        (
            CensusCbpAdapter(),
            "census-cbp-2023.json",
            "cbp-public-proxy-frame-batch",
        ),
    )
    batches = []
    for adapter, filename, batch_id in adapters:
        snapshot = _load(PUBLIC_FIXTURES / filename)
        batches.append(adapter.normalize(
            snapshot,
            {
                "batch_id": batch_id,
                "frame_request_id": request["request_id"],
            },
        ))
    return batches


def _bare_digest(document: object) -> str:
    return sha256_json(document).removeprefix("sha256:")


def _research_support_bindings(
    brief: dict[str, object],
    panel: dict[str, object],
) -> dict[str, str]:
    """Derive non-placeholder audit bindings from deterministic support docs."""

    v2_brief = _v2_projection(brief, brief=True)
    v2_panel = _v2_projection(panel, brief=False)
    evidence_ledger = {
        "brief_id": brief["brief_id"],
        "evidence_ids": sorted(
            source["evidence_id"] for source in brief["evidence_sources"]
        ),
    }
    finding_support = {
        "brief_id": brief["brief_id"],
        "findings": [{
            "finding_id": finding["finding_id"],
            "evidence_ids": sorted(finding["evidence_ids"]),
        } for finding in brief["findings"]],
    }
    synthesis_matrix = {
        "brief_id": brief["brief_id"],
        "segment_ids": sorted(
            segment["segment_id"] for segment in panel["segments"]
        ),
    }
    report_inputs = {
        "brief_sha256": _bare_digest(v2_brief),
        "panel_sha256": _bare_digest(v2_panel),
        "population_frame_result_sha256":
            brief["population_frame_result_sha256"],
    }
    first_pass = {
        "evidence_ledger_sha256": _bare_digest(evidence_ledger),
        "finding_support_sha256": _bare_digest(finding_support),
        "synthesis_matrix_sha256": _bare_digest(synthesis_matrix),
        "report_inputs_sha256": _bare_digest(report_inputs),
    }
    report_manifest = {
        "panel_id": panel["panel_id"],
        "panel_version": panel["version"],
        "inputs": first_pass,
    }
    return {
        **first_pass,
        "report_manifest_sha256": _bare_digest(report_manifest),
    }


def _build_construction_audit(
    *,
    brief: dict[str, object],
    panel: dict[str, object],
    frame: dict[str, object],
    composition: dict[str, object],
    validity: dict[str, object],
    authorized_handoff: dict[str, object] | None,
    auditor_run_id: str,
) -> dict[str, object]:
    v2_brief = _v2_projection(brief, brief=True)
    v2_panel = _v2_projection(panel, brief=False)
    support = _research_support_bindings(brief, panel)
    audit_binding = panel["audit_binding"]
    for key, expected in support.items():
        if audit_binding[key] != expected:
            raise AssertionError(
                f"panel audit binding {key} is not the exact support digest"
            )
    input_bindings = {
        "brief_sha256": _bare_digest(v2_brief),
        "panel_sha256": _bare_digest(v2_panel),
        "evidence_ledger_sha256": support["evidence_ledger_sha256"],
        "finding_support_sha256": support["finding_support_sha256"],
        "synthesis_matrix_sha256": support["synthesis_matrix_sha256"],
        "report_manifest_sha256": support["report_manifest_sha256"],
        "population_frame_result_sha256": _bare_digest(frame),
        "population_frame_sha256": (
            _bare_digest(frame)
            if frame["eligibility"] in {
                "eligible_tier_2",
                "eligible_tier_3",
            }
            else None
        ),
        "composition_plan_sha256": _bare_digest(composition),
        "validity_profile_sha256": _bare_digest(validity),
        "authorized_handoff_sha256": (
            None
            if authorized_handoff is None
            else _bare_digest(authorized_handoff)
        ),
    }
    first_evidence_id = brief["evidence_sources"][0]["evidence_id"]
    first_finding_id = brief["findings"][0]["finding_id"]
    first_profile_id = composition["profiles"][0]["profile_id"]
    check_paths = {
        "population_frame_traceability": [
            f"population_frame.cells[{frame['cells'][0]['cell_id']}]"
        ],
        "weight_semantics": [
            f"composition_plan.profiles[{first_profile_id}]"
        ],
        "authorized_handoff_traceability": (
            []
            if authorized_handoff is None
            else [
                "authorized_handoff.cohorts"
                f"[{brief['authorized_audience_import']['cohort_id']}]"
            ]
        ),
    }
    checks = []
    for check_id in (
        "approved_evidence_only",
        "finding_support_complete",
        "contradictions_preserved",
        "segment_sufficiency",
        "profile_traceability",
        "inference_boundaries",
        "privacy_boundary",
        "count_semantics",
        "claim_tier",
        "population_frame_traceability",
        "weight_semantics",
        "authorized_handoff_traceability",
    ):
        checks.append({
            "check_id": check_id,
            "status": (
                "not_applicable"
                if (
                    check_id == "authorized_handoff_traceability"
                    and authorized_handoff is None
                )
                else "pass"
            ),
            "evidence_paths": check_paths.get(
                check_id,
                [f"brief.evidence_sources[{first_evidence_id}]"],
            ),
            "finding_ids": [first_finding_id],
            "profile_ids": [first_profile_id],
            "message": "Golden proof binds the exact Release B1 documents.",
        })
    return {
        "schema_version": "panel-construction-audit-v2",
        "applicability": "release_b1",
        "panel_id": panel["panel_id"],
        "panel_version": panel["version"],
        "auditor_run_id": auditor_run_id,
        "audited_at": "2026-07-24T17:00:00Z",
        "input_bindings": input_bindings,
        "checks": checks,
        "result": "pass",
        "limitations": [
            "The golden proof validates construction inputs, not ad outcomes."
        ],
    }


def _validate_complete_v3_chain(
    *,
    brief: dict[str, object],
    panel: dict[str, object],
    frame: dict[str, object],
    composition: dict[str, object],
    validity: dict[str, object],
    authorized_handoff: dict[str, object] | None,
    workflow_id: str,
    auditor_run_id: str,
) -> tuple[dict[str, object] | None, ...]:
    """Build real inline approval/audit docs and run the seven-doc gate."""

    v2_brief = _v2_projection(brief, brief=True)
    v2_panel = _v2_projection(panel, brief=False)
    brief_sha256 = _bare_digest(v2_brief)
    panel_sha256 = _bare_digest(v2_panel)
    expected_scoped_approvals = {
        "evidence-synthesis": "sha256:" + brief_sha256,
        "panel-construction": "sha256:" + panel_sha256,
    }
    for approval in brief["scoped_approvals"]:
        expected = expected_scoped_approvals[approval["scope"]]
        if approval["status"] != "approved":
            raise AssertionError("every scoped golden approval must be approved")
        if approval["target_sha256"] != expected:
            raise AssertionError(
                "scoped golden approval does not bind the exact v2 document"
            )

    audit = _build_construction_audit(
        brief=brief,
        panel=panel,
        frame=frame,
        composition=composition,
        validity=validity,
        authorized_handoff=authorized_handoff,
        auditor_run_id=auditor_run_id,
    )
    audit_sha256 = _bare_digest(audit)
    if panel["audit_binding"]["audit_sha256"] != audit_sha256:
        raise AssertionError(
            "panel audit binding is not the exact audit digest: "
            + audit_sha256
        )
    support_digests = {
        panel["audit_binding"][key]
        for key in (
            "report_inputs_sha256",
            "evidence_ledger_sha256",
            "finding_support_sha256",
            "synthesis_matrix_sha256",
            "report_manifest_sha256",
        )
    }
    if len(support_digests) != 5:
        raise AssertionError("support bindings must be distinct real digests")
    workflow_state = {
        "schema_version": "panel-workflow-state-v1",
        "workflow_id": workflow_id,
        "panel_id": panel["panel_id"],
        "panel_version": panel["version"],
        "state": "approved",
        "updated_at": "2026-07-24T17:15:00Z",
        "approvals": [{
            "scope": "evidence_synthesis",
            "status": "approved",
            "approved_by": "golden-reviewer",
            "approved_at": "2026-07-24T17:10:00Z",
            "target_sha256": brief_sha256,
            "note": "Exact v2 brief projection approved.",
        }, {
            "scope": "panel_construction",
            "status": "approved",
            "approved_by": "golden-reviewer",
            "approved_at": "2026-07-24T17:12:00Z",
            "target_sha256": panel_sha256,
            "note": "Exact v2 panel projection approved.",
        }],
        "bindings": {
            "brief_sha256": brief_sha256,
            "panel_sha256": panel_sha256,
            "report_inputs_sha256":
                panel["audit_binding"]["report_inputs_sha256"],
            "audit_sha256": audit_sha256,
            "package_sha256": None,
        },
    }
    return validate_audience_research_v3(
        brief,
        panel,
        frame=frame,
        composition=composition,
        validity=validity,
        workflow_state=workflow_state,
        construction_audit=audit,
    )


_B2_FORBIDDEN_KEYS = {
    "package_manifest", "package_manifests",
    "package_archive", "package_archives", "archive", "archives",
    "resolver_envelope", "resolver_envelopes",
    "resolver_output", "resolver_outputs",
    "resolution_envelope", "resolution_envelopes",
    "resolution_output", "resolution_outputs",
}
_B2_FORBIDDEN_KEY_TOKENS = {
    "quota", "quotas", "job", "jobs", "assignment", "assignments",
    "result", "results", "score", "scores", "rank", "ranks",
    "dashboard", "dashboards", "package", "packages", "archive", "archives",
    "resolver", "resolvers", "resolution", "resolutions",
}
_B1_RESULT_BINDING_KEYS = {
    "frame_result_sha256",
    "population_frame_result_sha256",
}
_B1_RESULT_BINDING_LOCATIONS = {
    (
        "audience-research-brief-v3",
        (),
        "population_frame_result_sha256",
        True,
    ),
    (
        "saved-audience-panel-v3",
        (),
        "population_frame_result_sha256",
        True,
    ),
    (
        "panel-composition-plan-v1",
        ("frame_binding",),
        "frame_result_sha256",
        True,
    ),
    (
        "panel-validity-profile-v1",
        ("source_bindings",),
        "frame_result_sha256",
        True,
    ),
    (
        "panel-construction-audit-v2",
        ("input_bindings",),
        "population_frame_result_sha256",
        False,
    ),
}


def _assert_b1_only_documents(documents: object) -> None:
    """Recursively reject Release B2 document structures and archives."""

    def is_canonical_digest(value: object, *, prefixed: bool) -> bool:
        if not isinstance(value, str):
            return False
        digest = value
        if prefixed:
            if not digest.startswith("sha256:"):
                return False
            digest = digest.removeprefix("sha256:")
        return (
            len(digest) == 64
            and all(character in "0123456789abcdef" for character in digest)
        )

    def walk(value: object, path: tuple[str, ...], root_schema: object) -> None:
        if isinstance(value, dict):
            for raw_key, child in value.items():
                normalized = str(raw_key).lower().replace("-", "_")
                key_tokens = set(normalized.split("_"))
                allowed_audit_result = (
                    raw_key == "result"
                    and not path
                    and root_schema == "panel-construction-audit-v2"
                    and isinstance(child, str)
                    and child in {"pass", "fail"}
                )
                result_binding_location = next(
                    (
                        location
                        for location in _B1_RESULT_BINDING_LOCATIONS
                        if location[:3] == (root_schema, path, raw_key)
                    ),
                    None,
                )
                allowed_result_binding = (
                    raw_key in _B1_RESULT_BINDING_KEYS
                    and result_binding_location is not None
                    and is_canonical_digest(
                        child,
                        prefixed=result_binding_location[3],
                    )
                )
                allowed_package_status = (
                    raw_key == "package_status"
                    and child == "unpackaged"
                    and not path
                    and root_schema == "saved-audience-panel-v3"
                )
                allowed_null_package_binding = (
                    raw_key == "package_sha256"
                    and path == ("bindings",)
                    and root_schema == "panel-workflow-state-v1"
                    and child is None
                )
                forbidden = (
                    normalized in _B2_FORBIDDEN_KEYS
                    or bool(key_tokens & _B2_FORBIDDEN_KEY_TOKENS)
                )
                if (
                    forbidden
                    and not allowed_audit_result
                    and not allowed_result_binding
                    and not allowed_package_status
                    and not allowed_null_package_binding
                ):
                    raise AssertionError(
                        "Release B2 key crossed the B1 boundary at "
                        + ".".join((*path, str(raw_key)))
                    )
                walk(child, (*path, str(raw_key)), root_schema)
        elif isinstance(value, list):
            for index, child in enumerate(value):
                walk(child, (*path, f"[{index}]"), root_schema)
        elif isinstance(value, str) and value.lower().endswith(".zip"):
            raise AssertionError(
                "Release B2 archive crossed the B1 boundary at "
                + ".".join(path)
            )

    if isinstance(documents, tuple):
        for document in documents:
            root_schema = (
                document.get("schema_version")
                if isinstance(document, dict)
                else None
            )
            walk(document, (), root_schema)
    else:
        root_schema = (
            documents.get("schema_version")
            if isinstance(documents, dict)
            else None
        )
        walk(documents, (), root_schema)


class PopulationCoreGoldenPathTests(unittest.TestCase):
    maxDiff = None

    def test_public_proxy_golden_path(self) -> None:
        request = _load(PUBLIC_FIXTURES / "frame-request.json")
        frame = build_population_frame(
            frame_request=request,
            observation_batches=_public_batches(request),
            built_at=BUILT_AT,
        )
        self.assertEqual(
            _load(PUBLIC_FIXTURES / "expected-population-frame.json"),
            frame,
        )

        evidence = _load(PUBLIC_FIXTURES / "overlay-evidence.json")
        composition = build_composition_plan(
            population_frame=frame,
            structural_findings=evidence["structural_findings"],
            overlay_findings=evidence["overlay_findings"],
            supported_profile_specs=evidence["profile_specs"],
            requested_tier="tier_2",
            evidence_basis="public",
            plan_id="marketing-leader-proxy-composition",
            plan_version="1.0.0",
            built_at=COMPOSED_AT,
        )
        expected_composition = _load(
            PUBLIC_FIXTURES / "expected-composition-plan.json"
        )
        for profile in expected_composition["profiles"]:
            profile["support_status"] = "supported"
        expected_composition["allocation_constraints"][0] = (
            "Preserve every explicit materialized profile signature."
        )
        self.assertEqual(expected_composition, composition)

        units = {
            (item["unit"], item["denominator"])
            for item in frame["units"]
        }
        self.assertEqual({
            ("persons", "employed-persons-excluding-self-employed"),
            ("firms", "employer-firms"),
            ("establishments", "employer-establishments"),
        }, units)
        self.assertEqual(len(units), len({
            item["partition_id"] for item in frame["units"]
        }))
        self.assertIn("public proxy", frame["claim_boundary"].lower())
        self.assertIn(
            "These proxies do not represent the full commercial target audience.",
            frame["claim_boundary"],
        )
        self.assertEqual(
            {"observed", "modeled", "missing"},
            {cell["status"] for cell in frame["cells"]},
        )
        missing_joint = next(
            joint
            for joint in frame["joints"]
            if joint["dimensions"] == ["employment-status", "geography"]
        )
        self.assertEqual([], missing_joint["cell_ids"])
        self.assertEqual(
            "The required critical joint is not available from the selected "
            "source observations.",
            missing_joint["missing_reason"],
        )
        self.assertEqual("experimental", frame["eligibility"])
        self.assertEqual(
            "missing-critical-joint:"
            "persons-employed-persons-excluding-self-employed:"
            "employment-status-geography;modeled-share-above-threshold",
            frame["downgrade_reason"],
        )

        tier_two_request = deepcopy(request)
        tier_two_request.update({
            "request_id": "public-firm-tier-2-frame",
            "target_audience": "U.S. professional-services employer firms",
            "target_unit": "firms",
            "proxy_universes": [request["proxy_universes"][1]],
            "required_dimensions": [
                "enterprise-size",
                "geography",
                "industry",
            ],
            "required_joints": [[
                "enterprise-size",
                "geography",
                "industry",
            ]],
        })
        tier_two_batch = CensusSusbAdapter().normalize(
            _load(PUBLIC_FIXTURES / "census-susb-2022.json"),
            {
                "batch_id": "susb-tier-2-frame-batch",
                "frame_request_id": tier_two_request["request_id"],
            },
        )
        tier_two_frame = build_population_frame(
            frame_request=tier_two_request,
            observation_batches=[tier_two_batch],
            built_at=BUILT_AT,
        )
        self.assertEqual("eligible_tier_2", tier_two_frame["eligibility"])
        self.assertEqual("", tier_two_frame["downgrade_reason"])
        self.assertEqual(
            {("firms", "employer-firms")},
            {
                (item["unit"], item["denominator"])
                for item in tier_two_frame["units"]
            },
        )
        self.assertEqual("tier_1", composition["achieved_tier"])
        self.assertEqual(
            ["no-eligible-population-frame"],
            composition["tier_reason_codes"],
        )

        brief = _load(PUBLIC_FIXTURES / "expected-v3-brief.json")
        panel = _load(PUBLIC_FIXTURES / "expected-v3-panel.json")
        self.assertEqual(brief, validate_research_brief_v3(brief))
        self.assertEqual(panel, validate_saved_panel_v3(panel))
        self.assertEqual(sha256_json(frame), brief["population_frame_result_sha256"])
        self.assertIsNone(brief["population_frame_sha256"])
        self.assertEqual(
            sha256_json(composition),
            panel["composition_plan_sha256"],
        )
        self.assertEqual("tier_1", brief["panel_tier"])
        self.assertEqual("public", brief["evidence_basis"])
        self.assertEqual(brief["panel_tier"], panel["panel_tier"])
        self.assertEqual(brief["claim_boundary"], panel["claim_boundary"])
        provisional_validity = assess_population_validity(
            frame_request=request,
            population_frame=frame,
            overlay_evidence=[],
            outcome_feedback=[],
        )
        validity = finalize_validity_profile(
            provisional_validity=provisional_validity,
            population_frame=frame,
            composition_plan=composition,
            panel_id=panel["panel_id"],
            panel_tier="tier_1",
            evidence_basis="public",
            brief_sha256=sha256_json(
                _v2_projection(brief, brief=True)
            ),
            panel_projection_sha256=sha256_json(
                _v2_projection(panel, brief=False)
            ),
        )
        self.assertEqual(
            sha256_json(validity),
            panel["validity_profile_sha256"],
        )
        validated = _validate_complete_v3_chain(
            brief=brief,
            panel=panel,
            frame=frame,
            composition=composition,
            validity=validity,
            authorized_handoff=None,
            workflow_id="marketing-leader-public-proxy-build",
            auditor_run_id="public-proxy-construction-audit",
        )
        self.assertEqual(
            (brief, panel, frame, composition),
            validated[:4],
        )
        _assert_b1_only_documents(validated)
        for allowed_document in (
            {
                "schema_version": "panel-construction-audit-v2",
                "result": "pass",
            },
            {
                "schema_version": "panel-construction-audit-v2",
                "result": "fail",
            },
            {
                "schema_version": "saved-audience-panel-v3",
                "package_status": "unpackaged",
            },
            {
                "schema_version": "panel-workflow-state-v1",
                "bindings": {"package_sha256": None},
            },
        ):
            with self.subTest(allowed_document=allowed_document):
                _assert_b1_only_documents(allowed_document)
        for forbidden_document in (
            {"nested": [{"study-quota": 4}]},
            {"nested": {"quota_count": 4}},
            {"nested": {"job-ids": []}},
            {"nested": {"profile_assignment": {}}},
            {"nested": {"assignment-results": []}},
            {"nested": {"confidence_scores": []}},
            {"nested": {"candidate-rank": 1}},
            {"nested": {"dashboards": []}},
            {"nested": {"package-manifest": {}}},
            {"nested": {"package_archives": []}},
            {"nested": {"resolver-envelope": {}}},
            {"nested": {"resolution_outputs": []}},
            {"nested": {"package": {}}},
            {"nested": {"packages": []}},
            {"nested": {"package-output": {}}},
            {"nested": {"archive_path": "audience-panel.json"}},
            {"nested": {"v3-archive": {}}},
            {"nested": {"resolver_payload": {}}},
            {
                "schema_version": "saved-audience-panel-v3",
                "nested": {"package_status": "unpackaged"},
            },
            {
                "schema_version": "saved-audience-panel-v3",
                "package_status": "proposed",
            },
            {
                "schema_version": "panel-workflow-state-v1",
                "bindings": {"package_sha256": "f" * 64},
            },
            {
                "schema_version": "panel-construction-audit-v2",
                "Result": "pass",
            },
            {
                "schema_version": "panel-construction-audit-v2",
                "result": "passed",
            },
            {
                "schema_version": "saved-audience-panel-v3",
                "package-status": "unpackaged",
            },
            {
                "schema_version": "panel-workflow-state-v1",
                "bindings": {"package-sha256": None},
            },
            {
                "schema_version": "untrusted-nested-schema-v1",
                "bindings": {
                    "schema_version": "panel-workflow-state-v1",
                    "package_sha256": None,
                },
            },
            {
                "schema_version": "panel-outcome-feedback-binding-v1",
                "panel_binding": {"package_status": "unpackaged"},
            },
            {
                "schema_version": "panel-calibration-refresh-proposal-v1",
                "panel_binding": {"package_status": "unpackaged"},
            },
            {
                "schema_version": "panel-construction-audit-v2",
                "nested": {"result": "pass"},
            },
            {
                "schema_version": "panel-workflow-state-v1",
                "nested": {"package_sha256": None},
            },
            {"nested": {"frame-result-sha256": "f" * 64}},
            {"nested": {"archive_path": "audience-package-v3.zip"}},
        ):
            with self.subTest(forbidden_document=forbidden_document):
                with self.assertRaisesRegex(AssertionError, "Release B2"):
                    _assert_b1_only_documents(forbidden_document)

    def test_authorized_marketplace_shapes_converge(self) -> None:
        from openpyxl import Workbook

        helper = authorized_transform_test_support.AuthorizedAudienceTransformTests(
            methodName=(
                "test_flat_nested_linked_and_generated_xlsx_converge_"
                "to_identical_documents"
            )
        )
        helper.setUp()
        try:
            with tempfile.TemporaryDirectory() as temporary:
                root = Path(temporary)
                shape_sources: list[tuple[str, list[Path]]] = []

                flat_root = root / "flat"
                flat_root.mkdir()
                flat = flat_root / "flat-structural.csv"
                flat.write_bytes(
                    (
                        AUTHORIZED_INPUTS
                        / "source-shapes"
                        / "flat-structural.csv"
                    ).read_bytes()
                )
                shape_sources.append(("flat", [flat]))

                nested_root = root / "nested"
                nested_root.mkdir()
                nested = nested_root / "nested-export.json"
                nested.write_bytes(
                    (
                        AUTHORIZED_INPUTS
                        / "source-shapes"
                        / "nested-export.json"
                    ).read_bytes()
                )
                shape_sources.append(("nested", [nested]))

                linked_root = root / "linked"
                linked_root.mkdir()
                linked = []
                for filename in (
                    "linked-cohorts.csv",
                    "linked-distributions.csv",
                ):
                    destination = linked_root / filename
                    destination.write_bytes(
                        (
                            AUTHORIZED_INPUTS / "source-shapes" / filename
                        ).read_bytes()
                    )
                    linked.append(destination)
                shape_sources.append(("linked", linked))

                xlsx_root = root / "xlsx"
                xlsx_root.mkdir()
                xlsx = xlsx_root / "cohort.xlsx"
                workbook = Workbook()
                worksheet = workbook.active
                worksheet.title = "Cohort Summary"
                worksheet.append([
                    "metric",
                    "operations_leaders",
                    "finance_leaders",
                ])
                worksheet.append(["respondent_count", 420, 680])
                worksheet.append(["share", 0.38, 0.62])
                workbook.save(xlsx)
                shape_sources.append(("xlsx", [xlsx]))

                output_bytes = []
                normalized_reports = []
                decisions = []
                flat_handoff = None
                flat_output = None
                for shape, primary_paths in shape_sources:
                    bundle = primary_paths[0].parent
                    social_path = bundle / "topic-social-evidence.json"
                    social_path.write_bytes(
                        canonical_json_bytes(_topic_social_evidence())
                    )
                    outcomes_path = bundle / "aggregate-outcomes.json"
                    outcomes_path.write_bytes(
                        (
                            AUTHORIZED_INPUTS / "aggregate-outcomes.json"
                        ).read_bytes()
                    )
                    profile = profile_authorized_bundle(
                        [*primary_paths, social_path, outcomes_path],
                        profile_id="fictional-marketplace-cohort",
                        profile_version="1.0.0",
                        profiled_at=AUTHORIZED_PROFILED_AT,
                    )
                    decisions.append(profile["decision"]["status"])
                    mapping = _apply_authorized_golden_metadata(
                        helper._mapping_for_shape(shape, profile)
                    )
                    output = root / f"output-{shape}"
                    handoff = transform_authorized_bundle(
                        source_profile=profile,
                        mapping=mapping,
                        input_root=bundle,
                        output_dir=output,
                        transformer_version=AUTHORIZED_TRANSFORMER_VERSION,
                    )
                    if shape == "flat":
                        flat_handoff = deepcopy(handoff)
                        flat_output = output
                    self.assertEqual(
                        handoff,
                        validate_authorized_handoff(
                            handoff,
                            output_root=output,
                        ),
                    )
                    output_bytes.append({
                        name: (output / name).read_bytes()
                        for name in (
                            "frame-observations-0001.json",
                            "structured-evidence-0001.json",
                            "social-observations-0001.json",
                            "outcome-feedback-0001.json",
                            "outcome-feedback-0002.json",
                        )
                    })
                    report = _load(output / "transformation-report.json")
                    self.assertEqual(
                        {
                            "consequences": [],
                            "coverage": (
                                "all_selected_rows_and_fields_preserved"
                            ),
                        },
                        report["loss_summary"],
                    )
                    shape_provenance = {
                        "source_profile",
                        "mapping",
                        "input_hashes",
                        "source_reads",
                        "operation_log",
                        "field_changes",
                    }
                    normalized_reports.append({
                        key: value
                        for key, value in report.items()
                        if key not in shape_provenance
                    })

                self.assertEqual(["ready_for_mapping"] * 4, decisions)
                self.assertEqual([output_bytes[0]] * 4, output_bytes)
                self.assertEqual(
                    [normalized_reports[0]] * 4,
                    normalized_reports,
                )
                assert flat_handoff is not None
                assert flat_output is not None
                self.assertEqual(
                    _load(AUTHORIZED_FIXTURES / "expected-handoff.json"),
                    flat_handoff,
                )

                observation_batch = AuthorizedHandoffAdapter().normalize({
                    "handoff": flat_handoff,
                    "output_root": str(flat_output),
                })
                request = _authorized_request(calibration_factor=3.0)
                frame = build_population_frame(
                    frame_request=request,
                    observation_batches=[observation_batch],
                    built_at=BUILT_AT,
                )
                self.assertEqual(
                    _load(AUTHORIZED_FIXTURES / "expected-frame.json"),
                    frame,
                )
                self.assertEqual("eligible_tier_3", frame["eligibility"])
                self.assertEqual(
                    3.0,
                    next(
                        cell["calibration_factor"]
                        for cell in frame["cells"]
                        if cell["cell_id"] == "operations-leaders"
                    ),
                )

                structural, overlays, profiles = (
                    _authorized_composition_inputs()
                )
                composition = build_composition_plan(
                    population_frame=frame,
                    structural_findings=structural,
                    overlay_findings=overlays,
                    supported_profile_specs=profiles,
                    requested_tier="tier_3",
                    evidence_basis="first_party_aggregate",
                    plan_id="fictional-marketplace-composition",
                    plan_version="1.0.0",
                    built_at="2026-07-24T14:00:00Z",
                )
                self.assertEqual(
                    _load(
                        AUTHORIZED_FIXTURES / "expected-composition.json"
                    ),
                    composition,
                )
                self.assertEqual("tier_3", composition["achieved_tier"])
                self.assertEqual(
                    {
                        "operations",
                        "pricing",
                        "returns",
                        "customer-service",
                    },
                    {
                        binding["topic_id"]
                        for overlay in composition["overlay_hypotheses"]
                        for binding in overlay["topic_bindings"]
                    },
                )

                structured_evidence = _load(
                    flat_output / "structured-evidence-0001.json"
                )
                social_evidence = _load(
                    flat_output / "social-observations-0001.json"
                )
                frame_observations = _load(
                    flat_output / "frame-observations-0001.json"
                )
                semantic_inputs = _authorized_semantic_inputs()
                self.assertEqual(
                    sha256_json(semantic_inputs["frame"]),
                    frame_observations["raw_snapshot_sha256"],
                )
                self.assertEqual(
                    sha256_json(semantic_inputs["structured"]),
                    structured_evidence["input_sha256"],
                )
                self.assertEqual(
                    sha256_json(semantic_inputs["social"]),
                    social_evidence["input_sha256"],
                )
                structured_summaries = {
                    record["evidence_id"]: record["summary"]
                    for record in semantic_inputs["structured"]["records"]
                }
                for item in structured_evidence["items"]:
                    self.assertEqual(
                        structured_summaries[item["evidence_item_id"]],
                        item["content_summary"],
                    )
                    self.assertEqual(
                        sha256_text(item["content_summary"]),
                        item["content_sha256"],
                    )
                social_summaries = {
                    record["observation_id"]: record["summary"]
                    for record in semantic_inputs["social"]["records"]
                }
                for item in social_evidence["observations"]:
                    self.assertEqual(
                        social_summaries[item["observation_id"]],
                        item["text_excerpt"],
                    )
                    self.assertEqual(
                        sha256_text(item["text_excerpt"]),
                        item["content_sha256"],
                    )
                _assert_no_repeated_digest_placeholders((
                    frame_observations,
                    structured_evidence,
                    social_evidence,
                ))
                self.assertEqual(
                    {
                        "operations",
                        "pricing",
                        "returns",
                        "customer-service",
                    },
                    {
                        item["cluster_id"]
                        for item in social_evidence["observations"]
                    },
                )
                self.assertEqual(
                    {"authorized_aggregate_source"},
                    {
                        item["platform"]
                        for item in social_evidence["observations"]
                    },
                )
                self.assertTrue(all(
                    item["author_group_token"] is None
                    for item in social_evidence["observations"]
                ))
                self.assertEqual(
                    "Authorized aggregate source",
                    social_evidence["collection"]["provider"],
                )
                social_by_id = {
                    item["observation_id"]: item
                    for item in social_evidence["observations"]
                }
                composition_social_ids = {
                    evidence_id
                    for overlay in composition["overlay_hypotheses"]
                    for evidence_id in overlay["evidence_ids"]
                    if evidence_id.startswith("social-evidence-")
                }
                self.assertEqual(set(social_by_id), composition_social_ids)
                self.assertEqual(
                    {
                        "social-evidence-operations": "operations-leaders",
                        "social-evidence-customer-service":
                            "operations-leaders",
                        "social-evidence-pricing": "finance-leaders",
                        "social-evidence-returns": "finance-leaders",
                    },
                    {
                        evidence_id: profile["source_cell_ids"][0]
                        for profile in composition["profiles"]
                        for evidence_id in profile["support_evidence_ids"]
                        if evidence_id.startswith("social-evidence-")
                    },
                )
                frame_cell_ids = {
                    cell["cell_id"] for cell in frame["cells"]
                }
                self.assertTrue(
                    all(
                        profile["source_cell_ids"][0] in frame_cell_ids
                        for profile in composition["profiles"]
                        if any(
                            evidence_id.startswith("social-evidence-")
                            for evidence_id
                            in profile["support_evidence_ids"]
                        )
                    )
                )
                outcomes = [
                    _load(path)
                    for path in sorted(
                        flat_output.glob("outcome-feedback-*.json")
                    )
                ]
                brief = _load(
                    AUTHORIZED_FIXTURES / "expected-v3-brief.json"
                )
                panel = _load(
                    AUTHORIZED_FIXTURES / "expected-v3-panel.json"
                )
                self.assertEqual(brief, validate_research_brief_v3(brief))
                self.assertEqual(panel, validate_saved_panel_v3(panel))

                provisional = assess_population_validity(
                    frame_request=request,
                    population_frame=frame,
                    overlay_evidence=[
                        structured_evidence,
                        social_evidence,
                    ],
                    outcome_feedback=outcomes,
                )
                validity = finalize_validity_profile(
                    provisional_validity=provisional,
                    population_frame=frame,
                    composition_plan=composition,
                    panel_id=panel["panel_id"],
                    panel_tier="tier_3",
                    evidence_basis="first_party_aggregate",
                    brief_sha256=sha256_json(
                        _v2_projection(brief, brief=True)
                    ),
                    panel_projection_sha256=sha256_json(
                        _v2_projection(panel, brief=False)
                    ),
                )
                self.assertEqual(
                    _load(AUTHORIZED_FIXTURES / "expected-validity.json"),
                    validity,
                )
                self.assertEqual(
                    sha256_json(validity),
                    panel["validity_profile_sha256"],
                )
                self.assertEqual(
                    sha256_json(flat_handoff),
                    panel["authorized_handoff_sha256"],
                )
                validated = _validate_complete_v3_chain(
                    brief=brief,
                    panel=panel,
                    frame=frame,
                    composition=composition,
                    validity=validity,
                    authorized_handoff=flat_handoff,
                    workflow_id="fictional-marketplace-panel-build",
                    auditor_run_id=(
                        "fictional-marketplace-construction-audit"
                    ),
                )
                self.assertEqual(
                    (
                        brief,
                        panel,
                        frame,
                        composition,
                        validity,
                    ),
                    validated[:5],
                )
                _assert_b1_only_documents(validated)
                _assert_no_repeated_digest_placeholders(validated)

                panel_before = canonical_json_bytes(panel)
                frame_before = canonical_json_bytes(frame)
                composition_before = canonical_json_bytes(composition)
                feedback_binding = bind_outcome_feedback(
                    panel=panel,
                    feedback_documents=outcomes,
                    binding_id="marketplace-outcome-binding",
                    bound_at="2026-07-24T16:00:00Z",
                )
                self.assertEqual(
                    {"qualified-response-rate"},
                    {
                        metric["name"]
                        for metric in feedback_binding["metric_identities"]
                    },
                )
                self.assertEqual(
                    {"Seven days after exposure"},
                    {
                        metric["attribution_window"]
                        for metric in feedback_binding["metric_identities"]
                    },
                )
                proposal = propose_calibration_refresh(
                    panel=panel,
                    feedback_binding=feedback_binding,
                    proposal_id="marketplace-calibration-review",
                    proposed_at="2026-07-24T16:30:00Z",
                )
                self.assertFalse(proposal["executable"])
                self.assertEqual([], proposal["diff"]["operations"])
                self.assertIsNone(
                    proposal["diff"]["proposed_panel_sha256"]
                )
                self.assertEqual(panel_before, canonical_json_bytes(panel))
                self.assertEqual(frame_before, canonical_json_bytes(frame))
                self.assertEqual(
                    composition_before,
                    canonical_json_bytes(composition),
                )

                modeled_request = _authorized_request(
                    calibration_factor=3.0,
                    modeled_share=0.3000001,
                )
                modeled_batch = deepcopy(observation_batch)
                modeled_batch["frame_request_id"] = (
                    modeled_request["request_id"]
                )
                modeled_frame = build_population_frame(
                    frame_request=modeled_request,
                    observation_batches=[modeled_batch],
                    built_at=BUILT_AT,
                )
                self.assertEqual("experimental", modeled_frame["eligibility"])
                self.assertEqual(
                    "modeled-share-above-threshold",
                    modeled_frame["downgrade_reason"],
                )
                tier_one_structural = deepcopy(structural)
                for group, weight in zip(
                    tier_one_structural,
                    (0.38, 0.62),
                    strict=True,
                ):
                    group["cell_ids"] = []
                    group["planning_allocation"] = weight
                downgraded = build_composition_plan(
                    population_frame=modeled_frame,
                    structural_findings=tier_one_structural,
                    overlay_findings=overlays,
                    supported_profile_specs=profiles,
                    requested_tier="tier_3",
                    evidence_basis="first_party_aggregate",
                    plan_id="fictional-marketplace-modeled-downgrade",
                    plan_version="1.0.0",
                    built_at="2026-07-24T14:00:00Z",
                )
                self.assertEqual("tier_1", downgraded["achieved_tier"])
                self.assertEqual(
                    ["no-eligible-population-frame"],
                    downgraded["tier_reason_codes"],
                )
                with self.assertRaisesRegex(ValueError, "at most 3.0"):
                    build_population_frame(
                        frame_request=_authorized_request(
                            calibration_factor=3.0000001
                        ),
                        observation_batches=[observation_batch],
                        built_at=BUILT_AT,
                    )

                person_level = root / "person-level.csv"
                person_level.write_text(
                    "email,cohort\nperson@example.test,operations\n",
                    encoding="utf-8",
                )
                private_profile = profile_authorized_bundle(
                    [person_level],
                    profile_id="private-routing-proof",
                    profile_version="1.0.0",
                    profiled_at=AUTHORIZED_PROFILED_AT,
                )
                self.assertEqual(
                    "requires_private_aggregation",
                    private_profile["decision"]["status"],
                )
                self.assertNotIn(
                    "person@example.test",
                    canonical_json_bytes(private_profile).decode("utf-8"),
                )

                _assert_b1_only_documents((
                    flat_handoff,
                    frame,
                    composition,
                    validity,
                    brief,
                    panel,
                ))
        finally:
            helper.tearDown()
